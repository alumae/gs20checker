import asyncio
import base64
import json
import time
import tomllib
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app import create_app
from app.matching import AsyncMatcher
from app.processing import StreamBudget, process_workbook
from app.workbooks import extract_rows
from build import build
from tests.test_processing import FakeClient, all_negative, workbook


def test_vercel_project_declares_local_runtime_dependencies():
    root = Path(__file__).resolve().parents[1]
    project = tomllib.loads((root / "pyproject.toml").read_text())["project"]
    requirements = {
        line.strip()
        for line in (root / "requirements.txt").read_text().splitlines()
        if line.strip() and not line.lstrip().startswith("#")
    }
    assert project["name"] and project["version"]
    assert project["requires-python"] == ">=3.12"
    assert set(project["dependencies"]) == requirements


def test_vercel_limits_and_rendered_upload_configuration(monkeypatch, catalogue, tmp_path):
    path = tmp_path / "catalogue.json"
    path.write_text(json.dumps(catalogue.data))
    monkeypatch.setenv("VERCEL", "1")
    hosted = create_app({"TESTING": True, "GS_CATALOGUE": str(path)})
    assert hosted.config["MAX_CONTENT_LENGTH"] == 4_100_000
    assert hosted.config["GS_MAX_UPLOAD_BYTES"] == 4_000_000
    assert hosted.config["GS_MAX_RESPONSE_BYTES"] == 4_000_000
    assert hosted.config["GS_PROCESSING_SECONDS"] == 1740
    page = hosted.test_client().get("/").text
    assert 'data-max-bytes="4000000"' in page
    assert "29-minute processing limit" in page
    monkeypatch.delenv("VERCEL")
    local = create_app({"GS_CATALOGUE": str(path)})
    assert local.config["GS_MAX_UPLOAD_BYTES"] == 10 * 1024 * 1024
    assert local.config["GS_PROCESSING_SECONDS"] is None
    assert local.config["GS_MAX_RESPONSE_BYTES"] is None


def test_upload_limit_checked_before_excel_processing(application, client, headers, monkeypatch):
    application.config.update(GS_MAX_UPLOAD_BYTES=100, GS_UPLOAD_LIMIT_LABEL="100 bytes")

    def forbidden(_):
        raise AssertionError("Oversized input reached Excel parsing")

    monkeypatch.setattr("app.routes.extract_rows", forbidden)
    result = client.post(
        "/api/workbooks/process",
        data={"file": (BytesIO(b"x" * 101), "large.xlsx")},
        headers=headers,
    )
    assert result.status_code == 413 and "100 bytes" in result.json["error"]
    application.config["MAX_CONTENT_LENGTH"] = 50
    result = client.post(
        "/api/workbooks/process",
        data={"file": (BytesIO(b"x" * 101), "large.xlsx")},
        headers=headers,
    )
    assert result.status_code == 413 and "100 bytes" in result.json["error"]


def make_stream(catalogue, handler=all_negative, references=None, **limits):
    content = workbook(references or ["Reference"])
    rows, sheets = extract_rows(content)
    client = FakeClient(handler)
    engine = AsyncMatcher(catalogue, "test", client=client)
    stream = process_workbook(
        content, rows, sheets, catalogue, lambda: engine, heartbeat_seconds=0.001, **limits
    )
    return stream, client, rows


def test_deadline_cancels_pending_call_and_stops_later_batches(catalogue, monkeypatch):
    clock = [0]
    monkeypatch.setattr("app.processing.time", SimpleNamespace(monotonic=lambda: clock[0]))
    cancelled = []

    async def slow(_):
        try:
            clock[0] = 11
            await asyncio.sleep(3600)
        finally:
            cancelled.append(True)

    stream, client, rows = make_stream(
        catalogue, slow, references=[str(n) for n in range(20)], deadline=10
    )
    events = [json.loads(line) for line in stream]
    assert events[-1]["type"] == "error" and "time limit" in events[-1]["error"]
    assert not any(e["type"] in {"file_chunk", "complete"} for e in events)
    assert cancelled == [True] and client.closed and not rows and len(client.calls) == 1


def test_deadline_expired_before_start_never_calls_model(catalogue):
    stream, client, rows = make_stream(catalogue, deadline=time.monotonic() - 1)
    events = [json.loads(line) for line in stream]
    assert len(events) == 1 and events[0]["type"] == "error"
    assert not client.calls and not rows


def test_progress_bytes_are_budgeted_in_utf8(catalogue):
    stream, client, rows = make_stream(catalogue, references=["õ" * 3000], max_response_bytes=5000)
    lines = list(stream)
    events = [json.loads(line) for line in lines]
    assert events[-1]["type"] == "error" and "response size" in events[-1]["error"]
    assert len("".join(lines).encode("utf-8")) <= 5000
    assert not any(e["type"] == "file_chunk" for e in events)
    assert client.closed and not rows


@pytest.mark.parametrize("shortfall", [0, 1])
def test_download_preflight_counts_base64_framing_and_completion(catalogue, monkeypatch, shortfall):
    buffers = []

    def export(*_):
        output = BytesIO(b"x" * 100_001)  # Several chunks and a padded last chunk.
        buffers.append(output)
        return output

    monkeypatch.setattr("app.processing.export_workbook", export)
    unlimited, _, _ = make_stream(catalogue, references=["Jörgen"])
    expected = list(unlimited)
    limit = len("".join(expected).encode("utf-8")) + StreamBudget.ERROR_RESERVE - shortfall
    stream, client, rows = make_stream(catalogue, references=["Jörgen"], max_response_bytes=limit)
    lines = list(stream)
    events = [json.loads(line) for line in lines]
    assert len("".join(lines).encode("utf-8")) <= limit
    if shortfall:
        assert events[-1]["type"] == "error"
        assert not any(e["type"] == "file_chunk" for e in events)
    else:
        assert lines == expected and events[-1]["type"] == "complete"
    assert client.closed and not rows and all(b.closed for b in buffers)


def test_full_sample_fits_hosted_response_budget(application, client, headers):
    application.config.update(
        MAX_CONTENT_LENGTH=4_100_000,
        GS_MAX_UPLOAD_BYTES=4_000_000,
        GS_MAX_RESPONSE_BYTES=4_000_000,
        GS_PROCESSING_SECONDS=1740,
    )
    content = Path("sample_data/Publications.xlsx").read_bytes()
    response = client.post(
        "/api/workbooks/process",
        data={"file": (BytesIO(content), "sample.xlsx")},
        headers=headers,
    )
    assert response.status_code == 200 and len(response.data) <= 4_000_000
    events = [json.loads(line) for line in response.data.splitlines()]
    assert events[-1]["type"] == "complete" and events[-1]["total"] == 1161
    output = b"".join(base64.b64decode(e["data"]) for e in events if e["type"] == "file_chunk")
    book = load_workbook(BytesIO(output))
    assert book.active.max_row == 1162 and book.active.max_column == 63
    book.close()


def test_build_validates_catalogue_and_refreshes_cdn_assets(tmp_path, catalogue):
    (tmp_path / "data").mkdir()
    (tmp_path / "data/venues.json").write_text(json.dumps(catalogue.data))
    assets = tmp_path / "app/static"
    assets.mkdir(parents=True)
    (assets / "app.js").write_text("console.log('current')")
    (assets / "._app.js").write_text("metadata")
    build(tmp_path)
    public = tmp_path / "public/static"
    (public / "stale.js").write_text("old")
    build(tmp_path)
    assert sorted(p.name for p in public.iterdir()) == ["app.js"]
    assert (public / "app.js").read_text() == (assets / "app.js").read_text()

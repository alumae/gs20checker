import asyncio
import base64
import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.candidates import CandidateIndex
from app.catalogue import Catalogue
from app.matching import AsyncMatcher, Matcher, request_kwargs
from app.workbooks import save_workbook
from tests.test_processing import FakeClient, negative, reply, stream


@pytest.fixture(scope="module")
def real_catalogue():
    return Catalogue.load("data/venues.json")


@pytest.mark.parametrize(
    "name,expected",
    [
        (
            "International Journal of River Basin Management",
            "International Journal of River Basin Management",
        ),
        ("Computer Speech amp; Language", "Computer Speech & Language"),
        ("Computer Speech and Language", "Computer Speech & Language"),
        ("Biomass  Bioenergy", "Biomass and Bioenergy"),
        ("Materials Science and Engineering A", "Materials Science and Engineering: A"),
        ("Journal of Service Management", None),
        ("Materials Science", None),
        ("Organization", None),
        ("Music Perception", None),  # Subtitle/alias handling still belongs to the model.
    ],
)
def test_whole_journal_identity(real_catalogue, name, expected):
    venue_id = real_catalogue.candidates.exact_journal({"journal": name})
    assert (real_catalogue.venues[venue_id]["name"] if venue_id else None) == expected


def test_no_inference_from_titles_substrings_or_conference_tracks(real_catalogue):
    for fields in (
        {"title": "Computer Speech & Language"},
        {"reference": "A paper in Computer Speech & Language"},
        {"journal": "Workshop on Computer Speech & Language"},
        {"journal": "International Conference on Language Resources and Evaluation (LREC)"},
        {"journal": ""},
    ):
        assert real_catalogue.candidates.exact_journal(fields) is None
    index = CandidateIndex(
        {"a": {"name": "Testing & Research"}, "b": {"name": "Testing and Research"}}
    )
    assert index.exact_journal({"journal": "Testing Research"}) is None
    assert index.exact_journal({"journal": "Research Testing"}) is None


def test_river_basin_single_and_batch_never_call_model(real_catalogue):
    dataset = json.loads(Path("benchmarks/regressions/excel-identity.json").read_text())
    item = next(i for i in dataset["items"] if i["row"] == 997)
    # No configured API key: either path would fail if it consulted the model.
    single = Matcher(real_catalogue, "test")
    batch = AsyncMatcher(real_catalogue, "test")
    result = single.match(json.dumps(item["fields"]))
    assert result == asyncio.run(batch.match_batch([item]))[item["id"]]
    assert result["venue_id"] == item["expected_venue_id"]
    assert result["memberships"][0]["rank"] == 19
    assert single.requests == batch.requests == batch.verifications == 0


def test_exact_lookup_agrees_with_all_excel_regression_labels(real_catalogue):
    dataset = json.loads(Path("benchmarks/regressions/excel-identity.json").read_text())
    assert dataset["catalogue_version"] == real_catalogue.version
    resolved = []
    for item in dataset["items"]:
        venue_id = real_catalogue.candidates.exact_journal(item["fields"])
        if venue_id is not None:
            assert venue_id == item["expected_venue_id"]
            resolved.append(item["row"])
    assert resolved == [137, 157, 519, 804, 808, 997, 1010, 1120]


def test_excel_uses_only_full_references_and_exports_model_results(real_catalogue):
    book = Workbook()
    book.active.append(["Lühikirje", "Ajakiri / jätkuväljaanne"])
    reference = "Authors. Flood disasters. International Journal of River Basin Management."
    book.active.append([reference, "International Journal of River Basin Management"])
    book.active.append(["Unlisted article", "Journal of Service Management"])
    book.active.append([reference, "International Journal of Language & Communication Disorders"])
    book.active.append([None, "International Journal of River Basin Management"])
    venue_id = real_catalogue.candidates.exact_journal(
        {"journal": "International Journal of River Basin Management"}
    )

    async def handler(kwargs):
        items = json.loads(kwargs["input"][-1]["content"])
        results = []
        for item in items:
            result = negative(item["id"])
            if item["fields"]["reference"] == reference:
                result.update(
                    matched=True,
                    venue_id=venue_id,
                    venue_name=real_catalogue.venues[venue_id]["name"],
                    note="The reference identifies the journal.",
                )
            results.append(result)
        return reply(results)

    client = FakeClient(handler)
    events, engine, _ = stream(save_workbook(book).getvalue(), real_catalogue, client)
    events = [json.loads(e) for e in events]
    sent = json.loads(client.calls[0]["input"][-1]["content"])
    assert [item["fields"] for item in sent] == [
        {"reference": reference},
        {"reference": "Unlisted article"},
    ]
    assert engine.requests == 1 and engine.fallbacks == engine.verifications == 0
    data = b"".join(base64.b64decode(e["data"]) for e in events if e["type"] == "file_chunk")
    rows = list(load_workbook(BytesIO(data)).active.values)
    assert [r[2] for r in rows[1:]] == ["Yes", "No", "Yes", "No"]
    assert rows[1][3] == rows[3][3] == "International Journal of River Basin Management"
    assert rows[1][5] == "19"
    assert rows[3][1] == "International Journal of Language & Communication Disorders"
    assert rows[4][6] == "No usable publication reference."


def test_note_precedes_decision_and_name_precedes_id_in_schema(real_catalogue):
    single = request_kwargs(real_catalogue, "test", reference="A reference")
    batch = request_kwargs(real_catalogue, "test", items=[{"id": "a", "fields": {}}])
    assert list(single["text"]["format"]["schema"]["properties"]) == [
        "note",
        "matched",
        "venue_name",
        "venue_id",
    ]
    assert list(
        batch["text"]["format"]["schema"]["properties"]["results"]["items"]["properties"]
    ) == ["id", "note", "matched", "venue_name", "venue_id"]

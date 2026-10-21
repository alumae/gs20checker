import asyncio
import base64
import json
import sqlite3
import tempfile
import threading
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from types import SimpleNamespace

import httpx
import openai
import pytest
from openpyxl import Workbook, load_workbook

from app.matching import AsyncMatcher
from app.processing import process_workbook
from app.workbooks import extract_rows, save_workbook


def reply(results, batch=True):
    return SimpleNamespace(
        status="completed", output_text=json.dumps({"results": results} if batch else results)
    )


def negative(item_id=None):
    result = {"matched": False, "venue_id": None, "venue_name": None, "note": "Absent"}
    if item_id is not None:
        result["id"] = item_id
    return result


class FakeClient:
    def __init__(self, handler):
        self.handler = handler
        self.responses = self
        self.calls = []
        self.closed = False

    async def create(self, **kwargs):
        self.calls.append(kwargs)
        return await self.handler(kwargs)

    async def close(self):
        self.closed = True


def workbook(references):
    book = Workbook()
    book.active.append(["Lühikirje"])
    for reference in references:
        book.active.append([reference])
    return save_workbook(book).getvalue()


def stream(content, catalogue, client, heartbeat=0.001):
    rows, sheets = extract_rows(content)
    engine = AsyncMatcher(catalogue, "test", client=client)
    return (
        process_workbook(content, rows, sheets, catalogue, lambda: engine, heartbeat),
        engine,
        rows,
    )


async def all_negative(kwargs):
    items = json.loads(kwargs["input"][-1]["content"])
    return reply([negative(item["id"]) for item in reversed(items)])


def test_batch_validates_identity_and_retries_only_bad_item(catalogue):
    venue_id, venue = next(iter(catalogue.venues.items()))

    async def handler(kwargs):
        if len(client.calls) == 1:
            return reply(
                [
                    negative("good"),
                    {
                        "id": "bad",
                        "matched": True,
                        "venue_id": venue_id,
                        "venue_name": "Wrong name",
                        "note": "Looks plausible",
                    },
                ]
            )
        assert json.loads(kwargs["input"][-1]["content"]) == {"reference": "Bad reference"}
        return reply(
            {
                "matched": True,
                "venue_id": venue_id,
                "venue_name": venue["name"],
                "note": "Verified",
            },
            batch=False,
        )

    client = FakeClient(handler)
    matcher = AsyncMatcher(catalogue, "test", client=client)
    results = asyncio.run(
        matcher.match_batch(
            [
                {"id": "good", "fields": {"reference": "Good"}},
                {"id": "bad", "fields": {"reference": "Bad reference"}},
            ]
        )
    )
    assert results["good"]["status"] == "unmatched"
    assert results["bad"]["venue"] == venue["name"]
    assert matcher.requests == 2 and matcher.fallbacks == 1


@pytest.mark.parametrize("problem", ["missing", "duplicate", "unknown", "json", "incomplete"])
def test_bad_envelopes_fallback_individually(catalogue, problem):
    async def handler(kwargs):
        if len(client.calls) > 1:
            return reply(negative(), batch=False)
        results = [negative("a"), negative("b")]
        if problem == "missing":
            results.pop()
        if problem == "duplicate":
            results[1]["id"] = "a"
        if problem == "unknown":
            results[1]["id"] = "x"
        if problem == "json":
            return SimpleNamespace(status="completed", output_text="not JSON")
        if problem == "incomplete":
            return SimpleNamespace(status="incomplete", output_text="{}")
        return reply(results)

    client = FakeClient(handler)
    matcher = AsyncMatcher(catalogue, "test", client=client)
    results = asyncio.run(
        matcher.match_batch(
            [
                {"id": "a", "fields": {"reference": "A"}},
                {"id": "b", "fields": {"reference": "B"}},
            ]
        )
    )
    assert len(results) == 2 and all(r["status"] == "unmatched" for r in results.values())
    assert matcher.fallbacks == 2 and len(client.calls) == 3


def test_persistent_bad_output_is_error_not_nonmatch(catalogue):
    async def handler(kwargs):
        return SimpleNamespace(status="completed", output_text="{}")

    client = FakeClient(handler)
    matcher = AsyncMatcher(catalogue, "test", client=client)
    result = asyncio.run(matcher.match_batch([{"id": "a", "fields": {"reference": "A"}}]))
    assert result["a"]["status"] == "error"
    assert len(client.calls) == 2


def test_provider_error_never_triggers_individual_fallback(catalogue):
    async def handler(kwargs):
        raise openai.RateLimitError(
            "private diagnostics",
            body={"code": "insufficient_quota"},
            response=httpx.Response(429, request=httpx.Request("POST", "https://api.openai.com")),
        )

    client = FakeClient(handler)
    response, engine, rows = stream(workbook(["A", "B"]), catalogue, client)
    events = [json.loads(line) for line in response]
    assert events[-1]["type"] == "error"
    assert "billing" in events[-1]["error"]
    assert "private" not in str(events)
    assert not any(e["type"] in {"file_chunk", "complete"} for e in events)
    assert engine.fallbacks == 0 and len(client.calls) == 1
    assert client.closed and not rows


def test_deduplication_is_upload_local_and_results_keep_row_order(catalogue):
    content = workbook(["A", "A"] + [str(i) for i in range(8)] + ["A", "B"])
    for _ in range(2):
        client = FakeClient(all_negative)
        response, engine, rows = stream(content, catalogue, client)
        events = [json.loads(line) for line in response]
        assert engine.requests == 2
        inputs = [json.loads(call["input"][-1]["content"]) for call in client.calls]
        assert [len(items) for items in inputs] == [9, 1]
        recent = [r for e in events if e["type"] == "progress" for r in e["recent"]]
        assert [r["row"] for r in recent] == list(range(2, 14))
        assert events[-1]["counts"]["unmatched"] == 12
        data = b"".join(base64.b64decode(e["data"]) for e in events if e["type"] == "file_chunk")
        assert load_workbook(BytesIO(data)).active["B13"].value == "No"
        assert client.closed and not rows


def test_close_during_api_call_cancels_and_releases_upload(catalogue):
    cancelled = []

    async def handler(kwargs):
        try:
            await asyncio.sleep(3600)
        finally:
            cancelled.append(True)

    client = FakeClient(handler)
    response, engine, rows = stream(workbook([str(i) for i in range(20)]), catalogue, client)
    assert json.loads(next(response))["type"] == "started"
    assert json.loads(next(response))["type"] == "heartbeat"
    response.close()
    assert cancelled == [True]
    assert client.closed and not rows
    assert len(client.calls) == 1


def test_two_uploads_do_not_serialize_or_share_cancellation(catalogue):
    barrier = threading.Barrier(2)
    content = workbook(["A"])

    def run(cancel):
        async def handler(kwargs):
            # A barrier would time out if a global lock serialized the uploads.
            barrier.wait(timeout=3)
            await asyncio.sleep(0.03)
            return await all_negative(kwargs)

        client = FakeClient(handler)
        response, _, _ = stream(content, catalogue, client)
        next(response)
        first = json.loads(next(response))
        if cancel:
            response.close()
            return client.closed
        remaining = [first] + [json.loads(line) for line in response]
        return remaining[-1]["type"] == "complete" and client.closed

    with ThreadPoolExecutor(max_workers=2) as pool:
        a, b = pool.submit(run, True), pool.submit(run, False)
        assert a.result(timeout=5) and b.result(timeout=5)


def test_upload_and_export_never_create_database_or_temp_files(client, headers, monkeypatch):
    # ZIP-stored random strings make a genuine workbook above the 500KB spooling threshold.
    import random

    book = Workbook()
    book.active.append(["Lühikirje", "Preserved data"])
    rng = random.Random(42)
    for i in range(35):
        book.active.append([str(i), "".join(rng.choices("0123456789abcdef", k=30000))])
    data = save_workbook(book).getvalue()
    assert len(data) > 500000

    def forbidden(*args, **kwargs):
        raise AssertionError("Disk-backed processing is forbidden")

    monkeypatch.setattr(tempfile, "NamedTemporaryFile", forbidden)
    monkeypatch.setattr(tempfile, "TemporaryFile", forbidden)
    monkeypatch.setattr(tempfile, "SpooledTemporaryFile", forbidden)
    monkeypatch.setattr("werkzeug.formparser.SpooledTemporaryFile", forbidden)
    monkeypatch.setattr("werkzeug.formparser.TemporaryFile", forbidden, raising=False)
    monkeypatch.setattr("openpyxl.worksheet._writer.create_temporary_file", forbidden)
    monkeypatch.setattr(sqlite3, "connect", forbidden)
    response = client.post(
        "/api/workbooks/process", data={"file": (BytesIO(data), "ETIS.xlsx")}, headers=headers
    )
    events = [json.loads(line) for line in response.data.splitlines()]
    assert events[-1]["type"] == "complete"
    assert events[-1]["counts"]["pending"] == 0


def test_close_before_first_model_request_does_not_start_work(catalogue):
    client = FakeClient(all_negative)
    response, _, rows = stream(workbook(["A"]), catalogue, client)
    next(response)
    response.close()
    assert not rows and not client.calls


def test_cleanup_failure_still_releases_upload(catalogue, caplog):
    class BrokenCloseClient(FakeClient):
        async def close(self):
            raise RuntimeError("private diagnostics")

    client = BrokenCloseClient(all_negative)
    response, _, rows = stream(workbook(["A"]), catalogue, client)
    events = [json.loads(line) for line in response]
    assert events[-1]["type"] == "complete"
    assert not rows
    assert "RuntimeError" in caplog.text
    assert "private diagnostics" not in caplog.text

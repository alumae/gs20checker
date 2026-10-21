from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.workbooks import COLUMNS, export_workbook, extract_rows, save_workbook


def small_workbook():
    book = Workbook()
    sheet = book.active
    sheet.append(["Lühikirje", "Pealkiri", "Aasta", "Other"])
    sheet.append(["A. Test Journal", "Title", 2025, "=1+1"])
    sheet.append([None, None, None, None])
    sheet.append([None, None, 2024, "No reference"])
    book.create_sheet("Notes").append(["Keep this sheet"])
    return save_workbook(book).getvalue()


def test_extract_and_export(catalogue):
    content = small_workbook()
    rows, sheets = extract_rows(content)
    assert len(rows) == 2
    assert rows[1]["fields"] == {}
    result = catalogue.result(True, next(iter(catalogue.venues)), "=unsafe model formula")
    rows[0]["result"] = result
    output = export_workbook(content, {"rows": rows, "sheets": sheets, "snapshot": "test"})
    book = load_workbook(output)
    assert book.active["D2"].value == "=1+1"
    assert book.active["D2"].data_type == "f"
    assert book.active["I2"].value == "=unsafe model formula"
    assert book.active["I2"].data_type == "s"
    assert book.active["E3"].value is None
    assert book.active["E4"].value == "Pending"
    assert book["Notes"]["A1"].value == "Keep this sheet"
    assert len(book.active["G2"].value.splitlines()) == len(book.active["H2"].value.splitlines())


def test_sample_workbook_full_roundtrip(catalogue, tmp_path):
    content = Path("sample_data/Publications.xlsx").read_bytes()
    rows, sheets = extract_rows(content)
    assert len(rows) == 1161
    for row in rows:
        row["result"] = catalogue.result(False, None, "Synthetic test response.")
    before = load_workbook(BytesIO(content))
    after = load_workbook(
        export_workbook(content, {"rows": rows, "sheets": sheets, "snapshot": "test"})
    )
    assert before.sheetnames == after.sheetnames
    for old in before:
        new = after[old.title]
        assert old.max_row == new.max_row
        for row in old:
            for cell in row:
                other = new[cell.coordinate]
                assert cell.value == other.value
                assert cell.data_type == other.data_type
                assert cell.number_format == other.number_format
        assert [new.cell(1, old.max_column + i).value for i in range(1, 7)] == COLUMNS


def test_invalid_workbooks_and_limits():
    with pytest.raises(ValueError):
        extract_rows(b"not an excel file")
    with pytest.raises(ValueError, match="at most"):
        extract_rows(small_workbook(), max_rows=1)
    book = Workbook()
    book.active.append(["Unrelated header"])
    buf = save_workbook(book)
    with pytest.raises(ValueError, match="Lühikirje"):
        extract_rows(buf.getvalue())


def test_already_enriched_rejected(catalogue):
    content = small_workbook()
    rows, sheets = extract_rows(content)
    enriched = export_workbook(content, {"rows": rows, "sheets": sheets, "snapshot": "test"})
    with pytest.raises(ValueError, match="already has"):
        extract_rows(enriched.getvalue())

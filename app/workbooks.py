"""Read ETIS exports and append results without rebuilding the original sheets."""

from copy import copy
from io import BytesIO
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.writer.excel import ExcelWriter


class MemoryExcelWriter(ExcelWriter):
    """openpyxl 3.1 adapter: serialize worksheets without its default temp files."""

    def write_worksheet(self, ws):
        ws._drawing = SpreadsheetDrawing()
        ws._drawing.charts = ws._charts
        ws._drawing.images = ws._images
        with BytesIO() as buffer:
            writer = WorksheetWriter(ws, out=buffer)
            writer.write()
            ws._rels = writer._rels
            self._archive.writestr(ws.path[1:], buffer.getvalue())
        self.manifest.append(ws)


def save_workbook(workbook):
    output = BytesIO()
    try:
        with ZipFile(output, "w", ZIP_DEFLATED, allowZip64=True) as archive:
            MemoryExcelWriter(workbook, archive).save()
        output.seek(0)
        return output
    except BaseException:
        output.close()
        raise


COLUMNS = ["GS Match", "GS Venue", "GS Category", "GS Rank", "GS Note", "GS Snapshot"]
REFERENCE_HEADER = "Lühikirje"


def read_workbook(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("The expanded workbook exceeds 100 MB.")
            if len(archive.infolist()) > 10000:
                raise ValueError("The workbook contains too many internal files.")
        return load_workbook(BytesIO(content), data_only=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Upload a valid, unencrypted .xlsx workbook.") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not read this workbook as an Excel .xlsx file.") from exc


def extract_rows(content, max_rows=5000):
    workbook = read_workbook(content)
    rows, sheets = [], {}
    try:
        for sheet in workbook:
            if sheet.max_row > 50000 or sheet.max_column > 500:
                raise ValueError("Workbook dimensions exceed the supported ETIS layout.")
            headers = {str(c.value).strip(): c.column for c in sheet[1] if c.value is not None}
            if REFERENCE_HEADER not in headers:
                continue
            if any(name in headers for name in COLUMNS):
                raise ValueError(
                    "This workbook already has GS result columns. Upload the original."
                )
            sheets[sheet.title] = sheet.max_column
            for cells in sheet.iter_rows(min_row=2):
                if all(c.value is None for c in cells):
                    continue
                fields = {}
                cell = cells[headers[REFERENCE_HEADER] - 1]
                if cell.value is not None and cell.data_type != "f":
                    reference = str(cell.value)
                    if reference.strip():
                        # Send the complete citation unchanged, including BibTeX-like text.
                        # Other columns are preserved in the export but never used to match.
                        fields["reference"] = reference
                rows.append(
                    {
                        "sheet": sheet.title,
                        "row": cells[0].row,
                        "fields": fields,
                        "result": None,
                    }
                )
                if len(rows) > max_rows:
                    raise ValueError(f"Upload at most {max_rows:,} publication rows.")
        if not sheets:
            raise ValueError(
                "No ETIS publication sheet found: expected a Lühikirje header in row 1."
            )
        if not rows:
            raise ValueError("The workbook contains no publication rows.")
        return rows, sheets
    finally:
        workbook.close()


def export_workbook(content, job):
    workbook = read_workbook(content)
    try:
        for name, original_columns in job["sheets"].items():
            sheet = workbook[name]
            for offset, header in enumerate(COLUMNS, 1):
                cell = sheet.cell(1, original_columns + offset, header)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="176B5A")
                sheet.column_dimensions[cell.column_letter].width = [16, 42, 55, 12, 65, 30][
                    offset - 1
                ]
        for row in job["rows"]:
            result = row["result"] or {"status": "pending", "note": "Not processed yet."}
            members = result.get("memberships", [])
            values = [
                {"matched": "Yes", "unmatched": "No", "error": "Error", "pending": "Pending"}[
                    result["status"]
                ],
                result.get("venue", ""),
                "\n".join(m["category"] for m in members),
                "\n".join(str(m["rank"]) for m in members),
                result["note"],
                job["snapshot"],
            ]
            sheet = workbook[row["sheet"]]
            for offset, value in enumerate(values, 1):
                cell = sheet.cell(row["row"], job["sheets"][row["sheet"]] + offset)
                cell.value = value
                cell.data_type = "s"  # Model text and venue names are never Excel formulas.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row["row"] > 1:
                    cell.font = copy(sheet.cell(row["row"], 1).font)
        return save_workbook(workbook)
    finally:
        workbook.close()
